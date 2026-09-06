"""
exporter.py
-----------
Turn analysed results into downloadable reports, in several formats, shared by
both apps. The desktop app calls these for its "Save report" button; the web
app offers them as download buttons.

Formats:
  CSV   - always available (standard library only)
  JSON  - always available
  Excel - if openpyxl is installed
  Word  - if python-docx is installed
  PDF   - if reportlab is installed

Each function takes a list of result dicts (the same records used in the store)
and writes to the given path. Missing optional libraries are reported clearly
rather than crashing.
"""

import os
import csv
import json

FIELDS = ["name", "label", "confidence", "period", "duration", "depth",
          "depth_err", "snr", "planet_radius_earth", "timestamp"]


def to_csv(records, path):
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for r in records:
            w.writerow({k: r.get(k) for k in FIELDS})
    return path


def to_json(records, path):
    with open(path, "w") as f:
        json.dump(records, f, indent=2, default=str)
    return path


def to_excel(records, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except Exception:
        raise RuntimeError("Excel export needs openpyxl (pip install openpyxl)")
    wb = Workbook()
    ws = wb.active
    ws.title = "exopipe results"
    header_fill = PatternFill("solid", fgColor="1A2A4A")
    for c, name in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for r_i, rec in enumerate(records, 2):
        for c, name in enumerate(FIELDS, 1):
            ws.cell(row=r_i, column=c, value=rec.get(name))
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0)
                    for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 30)
    wb.save(path)
    return path


def to_word(records, path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except Exception:
        raise RuntimeError("Word export needs python-docx (pip install python-docx)")
    doc = Document()
    h = doc.add_heading("exopipe - Detection Report", level=0)
    doc.add_paragraph("AI-enabled exoplanet detection from TESS light curves")
    doc.add_paragraph("")

    for rec in records:
        doc.add_heading(str(rec.get("name", "target")), level=2)
        rows = [
            ("Class", "{}  (confidence {})".format(
                rec.get("label"), _pct(rec.get("confidence")))),
            ("Period (days)", _fmt(rec.get("period"))),
            ("Transit duration (days)", _fmt(rec.get("duration"))),
            ("Transit depth", _fmt(rec.get("depth"))),
            ("Depth uncertainty", _fmt(rec.get("depth_err"))),
            ("Transit SNR", _fmt(rec.get("snr"), 1)),
            ("Est. planet radius (Earth radii)",
             _fmt(rec.get("planet_radius_earth"), 1)),
            ("Analysed at", rec.get("timestamp", "")),
        ]
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for k, v in rows:
            cells = table.add_row().cells
            cells[0].text = k
            cells[1].text = str(v)
        # plot image if present
        plot = rec.get("plot")
        if plot and os.path.exists(plot):
            try:
                doc.add_picture(plot)
            except Exception:
                pass
        doc.add_paragraph("")
    doc.save(path)
    return path


def to_pdf(records, path):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.lib import colors
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                        Table, TableStyle, Image)
        from reportlab.lib.styles import getSampleStyleSheet
    except Exception:
        raise RuntimeError("PDF export needs reportlab (pip install reportlab)")

    styles = getSampleStyleSheet()
    story = [Paragraph("exopipe - Detection Report", styles["Title"]),
             Paragraph("AI-enabled exoplanet detection from TESS light curves",
                       styles["Normal"]),
             Spacer(1, 8)]
    for rec in records:
        story.append(Paragraph(str(rec.get("name", "target")),
                               styles["Heading2"]))
        data = [
            ["Class", "{} ({})".format(rec.get("label"),
                                       _pct(rec.get("confidence")))],
            ["Period (d)", _fmt(rec.get("period"))],
            ["Duration (d)", _fmt(rec.get("duration"))],
            ["Depth", _fmt(rec.get("depth"))],
            ["Depth err", _fmt(rec.get("depth_err"))],
            ["SNR", _fmt(rec.get("snr"), 1)],
            ["Planet radius (Earth)", _fmt(rec.get("planet_radius_earth"), 1)],
        ]
        t = Table(data, colWidths=[55 * mm, 90 * mm])
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2f8")),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(t)
        plot = rec.get("plot")
        if plot and os.path.exists(plot):
            try:
                story.append(Spacer(1, 4))
                story.append(Image(plot, width=160 * mm, height=44 * mm))
            except Exception:
                pass
        story.append(Spacer(1, 10))

    SimpleDocTemplate(path, pagesize=A4,
                      leftMargin=16 * mm, rightMargin=16 * mm,
                      topMargin=14 * mm, bottomMargin=14 * mm).build(story)
    return path


def _fmt(v, nd=4):
    if v is None:
        return "-"
    try:
        return "{:.{}f}".format(float(v), nd)
    except Exception:
        return str(v)


def _pct(v):
    try:
        return "{:.0%}".format(float(v))
    except Exception:
        return "-"


# what formats are available right now, for the UI to show/hide buttons
def available_formats():
    fmts = {"CSV": True, "JSON": True}
    for name, mod in (("Excel", "openpyxl"), ("Word", "docx"),
                      ("PDF", "reportlab")):
        try:
            __import__(mod)
            fmts[name] = True
        except Exception:
            fmts[name] = False
    return fmts


if __name__ == "__main__":
    recs = [
        {"name": "TIC 307210830", "label": "transit", "confidence": 0.82,
         "period": 4.6163, "duration": 0.12, "depth": 0.0035,
         "depth_err": 0.0001, "snr": 64.0, "planet_radius_earth": 6.4,
         "timestamp": "2026-01-01 12:00:00"},
    ]
    print("available formats:", available_formats())
    import tempfile
    d = tempfile.mkdtemp()
    print("csv  ->", to_csv(recs, os.path.join(d, "r.csv")))
    print("json ->", to_json(recs, os.path.join(d, "r.json")))
    for name, fn, ext in (("excel", to_excel, "xlsx"),
                          ("word", to_word, "docx"),
                          ("pdf", to_pdf, "pdf")):
        try:
            print(name, "->", fn(recs, os.path.join(d, "r." + ext)))
        except Exception as e:
            print(name, "skipped:", e)
